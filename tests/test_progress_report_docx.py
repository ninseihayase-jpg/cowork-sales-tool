"""社内PJ「進捗報告」docxエクスポート/インポート(2026-09-13)の回帰テスト。

ユーザー要望:「編集は、SFAで手修正でも可能だし、テンプレdocxをダウンロード→Uploadでも
可能」。build_progress_report_docx()で書き出したdocxを、そのままparse_progress_report_docx()
に通して往復できること（行の並びは_pr_ordered_rows()と一致させ、ラベル文言ではなく
行位置でセクションに対応させている）を中心に検証する。

一時DBのみ使用。本番DB(cowork_sfa.db)には一切触れない。
"""
from __future__ import annotations

import shutil
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path

import pytest

from cowork import sfa_db, webapp

BASIC_USER = "test_user"
BASIC_PASS = "test_pass_1234"


@pytest.fixture
def con():
    d = tempfile.mkdtemp(prefix="sfa_pr_docx_")
    path = str(Path(d) / "t.db")
    sfa_db.init_db(path)
    conn = sfa_db.connect(path)
    yield conn
    conn.close()
    shutil.rmtree(d, ignore_errors=True)


def _issue(con, issue="論点A"):
    acc = sfa_db.upsert_account(con, name="A社")
    did = sfa_db.upsert_deal(con, account_id=acc, deal_name="X", status="open")
    return sfa_db.upsert_deal_issue(con, deal_id=did, issue=issue)


def _seeded_report(con):
    iid = _issue(con)
    rep = sfa_db.open_progress_report_for_edit(con, iid, today="2026-09-06")
    sfa_db.update_progress_report(
        con, rep["id"], purpose_tag="decide",
        sections={
            "summary": "<div>順調です</div>",
            "progress": "<ul><li>タスクA完了</li><li>タスクB着手</li></ul>",
            "decision": "<div>予算をどうするか決めてほしい</div>",
            "risk_schedule": "<div>特になし</div>",
            "next_steps": "<ul><li>来週レビュー</li></ul>",
        })
    return iid, sfa_db.get_progress_report(con, rep["id"])


# ── HTML⇔行 変換ヘルパー ──

def test_html_to_lines_converts_bullets_with_marker():
    lines = webapp._pr_html_to_lines("<ul><li>項目A</li><li>項目B</li></ul>")
    assert lines == ["・項目A", "・項目B"]


def test_html_to_lines_converts_plain_divs_without_marker():
    lines = webapp._pr_html_to_lines("<div>1行目</div><div>2行目</div>")
    assert lines == ["1行目", "2行目"]


def test_html_to_lines_empty_input_returns_empty_list():
    assert webapp._pr_html_to_lines("") == []
    assert webapp._pr_html_to_lines(None) == []


def test_lines_to_html_reconstructs_bullet_list():
    html_val = webapp._pr_lines_to_html(["・項目A", "・項目B"])
    assert html_val == "<ul><li>項目A</li><li>項目B</li></ul>"


def test_lines_to_html_reconstructs_plain_lines():
    html_val = webapp._pr_lines_to_html(["1行目", "2行目"])
    assert html_val == "<div>1行目</div><div>2行目</div>"


def test_lines_to_html_mixes_bullet_and_plain_correctly():
    html_val = webapp._pr_lines_to_html(["導入文", "・項目A", "・項目B", "まとめ"])
    assert html_val == "<div>導入文</div><ul><li>項目A</li><li>項目B</li></ul><div>まとめ</div>"


def test_lines_to_html_escapes_dangerous_text():
    html_val = webapp._pr_lines_to_html(["<script>alert(1)</script>"])
    assert "<script>" not in html_val
    assert "&lt;script&gt;" in html_val


# ── docx生成 ──

def test_build_progress_report_docx_returns_nonempty_bytes(con):
    _, report = _seeded_report(con)
    data = webapp.build_progress_report_docx("テストPJ", report)
    assert len(data) > 1000
    assert data[:2] == b"PK"  # docx(zip)のマジックナンバー


# ── 往復（エクスポート→インポート）──

def test_docx_roundtrip_preserves_bullet_sections(con):
    _, report = _seeded_report(con)
    data = webapp.build_progress_report_docx("テストPJ", report)
    parsed = webapp.parse_progress_report_docx(data)
    assert parsed is not None
    assert parsed["progress"] == "<ul><li>タスクA完了</li><li>タスクB着手</li></ul>"
    assert parsed["next_steps"] == "<ul><li>来週レビュー</li></ul>"


def test_docx_roundtrip_preserves_plain_sections(con):
    _, report = _seeded_report(con)
    data = webapp.build_progress_report_docx("テストPJ", report)
    parsed = webapp.parse_progress_report_docx(data)
    assert parsed["summary"] == "<div>順調です</div>"
    assert parsed["decision"] == "<div>予算をどうするか決めてほしい</div>"


def test_docx_roundtrip_includes_all_section_keys_even_if_empty(con):
    _, report = _seeded_report(con)
    data = webapp.build_progress_report_docx("テストPJ", report)
    parsed = webapp.parse_progress_report_docx(data)
    assert set(parsed.keys()) == set(sfa_db.PROGRESS_REPORT_SECTION_KEYS)
    assert parsed["risk_budget"] == ""  # 元々空だったセクション


def test_parse_progress_report_docx_rejects_non_docx_bytes():
    assert webapp.parse_progress_report_docx(b"not a docx file") is None


def test_parse_progress_report_docx_rejects_docx_without_table():
    from docx import Document
    from io import BytesIO
    doc = Document()
    doc.add_paragraph("テーブルなしの文書")
    buf = BytesIO()
    doc.save(buf)
    assert webapp.parse_progress_report_docx(buf.getvalue()) is None


# ── ルート ──

@pytest.fixture
def server(monkeypatch, tmp_path):
    db_path = str(tmp_path / "srv.db")
    sfa_db.init_db(db_path)
    monkeypatch.setattr(webapp, "SFA_BASIC_USER", BASIC_USER)
    monkeypatch.setattr(webapp, "SFA_BASIC_PASS", BASIC_PASS)
    handler_cls = webapp._make_handler(db_path, None)
    srv = ThreadingHTTPServer(("127.0.0.1", 0), handler_cls)
    port = srv.server_address[1]
    import threading
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    try:
        yield f"http://127.0.0.1:{port}", db_path
    finally:
        srv.shutdown()
        srv.server_close()
        t.join(timeout=5)


def _auth_header():
    import base64
    token = base64.b64encode(f"{BASIC_USER}:{BASIC_PASS}".encode()).decode()
    return {"Authorization": f"Basic {token}"}


def test_export_docx_route_returns_valid_docx(server):
    base, db_path = server
    con = sfa_db.connect(db_path)
    iid, report = _seeded_report(con)
    con.close()

    req = urllib.request.Request(
        f"{base}/deal-issue-progress-report/{report['id']}/export.docx",
        headers=_auth_header(), method="GET")
    resp = urllib.request.urlopen(req, timeout=10)
    assert resp.getcode() == 200
    assert resp.headers.get("Content-Type") == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    body = resp.read()
    assert body[:2] == b"PK"


def test_export_docx_route_missing_report_returns_404(server):
    base, _ = server
    req = urllib.request.Request(
        f"{base}/deal-issue-progress-report/999999/export.docx",
        headers=_auth_header(), method="GET")
    try:
        urllib.request.urlopen(req, timeout=10)
        assert False, "404を期待"
    except urllib.error.HTTPError as e:
        assert e.code == 404


def _multipart_body(boundary, fields, files):
    parts = []
    for k, v in fields.items():
        parts.append(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n")
    body = "".join(parts).encode()
    for k, (filename, content_bytes) in files.items():
        header = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; "
                  f"filename=\"{filename}\"\r\nContent-Type: application/octet-stream\r\n\r\n")
        body += header.encode() + content_bytes + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    return body


def test_import_docx_route_updates_latest_report(server):
    base, db_path = server
    con = sfa_db.connect(db_path)
    iid, report = _seeded_report(con)
    docx_bytes = webapp.build_progress_report_docx("テストPJ", report)
    con.close()

    # docxを一部書き換えて（実際のユーザー編集を模して）再アップロード
    from docx import Document
    from io import BytesIO
    doc = Document(BytesIO(docx_bytes))
    table = doc.tables[0]
    table.rows[0].cells[1].paragraphs[0].text = "アップロード経由で更新したサマリー"
    buf = BytesIO()
    doc.save(buf)
    modified_bytes = buf.getvalue()

    boundary = "----testboundary123"
    body = _multipart_body(boundary, {}, {"docx_file": ("report.docx", modified_bytes)})
    req = urllib.request.Request(
        f"{base}/deal-issue-progress-report/{report['id']}/import-docx",
        data=body,
        headers={**_auth_header(), "Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST")
    resp = urllib.request.urlopen(req, timeout=10)
    assert resp.getcode() == 200
    assert "docxを取り込みました" in resp.read().decode("utf-8")

    con2 = sfa_db.connect(db_path)
    fetched = sfa_db.get_progress_report(con2, report["id"])
    con2.close()
    assert "アップロード経由で更新したサマリー" in fetched["summary_html"]


def test_import_docx_route_rejects_past_version(server):
    base, db_path = server
    con = sfa_db.connect(db_path)
    iid = _issue(con)
    old = sfa_db.open_progress_report_for_edit(con, iid, today="2026-08-30")
    sfa_db.open_progress_report_for_edit(con, iid, today="2026-09-06")  # 新verへフォーク
    docx_bytes = webapp.build_progress_report_docx("テストPJ", old)
    con.close()

    boundary = "----testboundary456"
    body = _multipart_body(boundary, {}, {"docx_file": ("report.docx", docx_bytes)})
    req = urllib.request.Request(
        f"{base}/deal-issue-progress-report/{old['id']}/import-docx",
        data=body,
        headers={**_auth_header(), "Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST")
    resp = urllib.request.urlopen(req, timeout=10)
    assert "過去verには取り込めません" in resp.read().decode("utf-8")
