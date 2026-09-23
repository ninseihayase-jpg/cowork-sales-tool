"""取り込み原本（文字起こし）のdocxダウンロード機能の回帰テスト(2026-09-09)。

ユーザー要望: 「文字起こしをdocxでダウンロードできるように。個別商談の該当エリアにも
ボタン表示（本文、の横）」。貼付テキスト（原本ファイルが無いケース）でもdocx化して
ダウンロードできるようにした。
一時DBのみ使用。本番DB(cowork_sfa.db)には一切触れない。
"""
from __future__ import annotations

import base64
import shutil
import tempfile
import threading
import urllib.error
import urllib.parse
import urllib.request
from http.server import ThreadingHTTPServer
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from cowork import sfa_db, webapp

BASIC_USER = "test_user"
BASIC_PASS = "test_pass_1234"


@pytest.fixture
def con():
    d = tempfile.mkdtemp(prefix="sfa_intake_docx_")
    path = str(Path(d) / "t.db")
    sfa_db.init_db(path)
    conn = sfa_db.connect(path)
    yield conn
    conn.close()
    shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def server(tmp_path, monkeypatch):
    db_path = str(tmp_path / "srv.db")
    sfa_db.init_db(db_path)
    monkeypatch.setattr(webapp, "GOOGLE_CLIENT_ID", BASIC_USER)
    monkeypatch.setattr(webapp, "GOOGLE_CLIENT_SECRET", BASIC_PASS)
    handler_cls = webapp._make_handler(db_path, None)
    srv = ThreadingHTTPServer(("127.0.0.1", 0), handler_cls)
    port = srv.server_address[1]
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    try:
        yield f"http://127.0.0.1:{port}", db_path
    finally:
        srv.shutdown()
        srv.server_close()
        t.join(timeout=5)


def _auth_header():
    return {"Cookie": f"sfa_session={webapp._make_session_token()}"}


def _get(url):
    req = urllib.request.Request(url, headers=_auth_header(), method="GET")
    try:
        resp = urllib.request.urlopen(req, timeout=10)
        return resp.getcode(), resp.headers, resp.read()
    except urllib.error.HTTPError as e:
        return e.code, e.headers, e.read()


def _s(x):
    return x.decode() if isinstance(x, (bytes, bytearray)) else x


# ── build_intake_transcript_docx (単体) ──

def test_build_intake_transcript_docx_contains_transcript_lines():
    t = {"filename": "貼り付けテキスト", "created_at": "2026-08-28 06:33:00",
         "transcript": "発言者A: こんにちは。\n発言者B: よろしくお願いします。"}
    data = webapp.build_intake_transcript_docx(t)
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "発言者A: こんにちは。" in text
    assert "発言者B: よろしくお願いします。" in text


def test_build_intake_transcript_docx_handles_missing_transcript_gracefully():
    data = webapp.build_intake_transcript_docx({"filename": "空", "transcript": None})
    doc = Document(BytesIO(data))
    assert doc.paragraphs  # 例外を投げずに空文字列で1文書は生成できる


# ── _intake_originals_html: docxボタンの表示 ──

def test_intake_originals_html_has_docx_button(con):
    acc = sfa_db.upsert_account(con, name="A社")
    did = sfa_db.upsert_deal(con, account_id=acc, deal_name="X", stage="提案")
    tid = sfa_db.add_intake_transcript(con, kind="deal", entity_id=did, source="paste",
                                       transcript="テスト文字起こし本文")
    html = webapp._intake_originals_html(con, "deal", did, f"/deal/{did}")
    assert f"/intake-transcript/{tid}/docx" in html
    assert "docx" in html
    # 既存の「本文」「原本DL」ボタンと並んで表示され、後方互換を壊していないこと
    assert f"/intake-transcript/{tid}/view" in html


# ── HTTPルート ──

def test_intake_transcript_docx_route_returns_docx(server):
    base, db_path = server
    con = sfa_db.connect(db_path)
    try:
        acc = sfa_db.upsert_account(con, name="B社")
        did = sfa_db.upsert_deal(con, account_id=acc, deal_name="Y", stage="提案")
        tid = sfa_db.add_intake_transcript(con, kind="deal", entity_id=did, source="paste",
                                           filename="面談メモ", transcript="議事録の本文です。")
        con.commit()
    finally:
        con.close()

    code, headers, body = _get(f"{base}/intake-transcript/{tid}/docx")
    assert code == 200
    assert headers.get("Content-Type") == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    # 日本語ファイル名はRFC 6266のfilename*(UTF-8)形式でパーセントエンコードされる
    # （_content_dispositionの既存仕様。http.serverがヘッダをlatin-1エンコードするため）。
    import urllib.parse as _up
    assert f"filename*=UTF-8''{_up.quote('面談メモ.docx', safe='')}" in headers.get("Content-Disposition", "")
    doc = Document(BytesIO(body))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "議事録の本文です。" in text


def test_intake_transcript_docx_route_404_when_no_transcript(server):
    base, db_path = server
    con = sfa_db.connect(db_path)
    try:
        acc = sfa_db.upsert_account(con, name="C社")
        did = sfa_db.upsert_deal(con, account_id=acc, deal_name="Z", stage="提案")
        tid = sfa_db.add_intake_transcript(con, kind="deal", entity_id=did, source="paste",
                                           transcript="")
        con.commit()
    finally:
        con.close()

    code, _, _ = _get(f"{base}/intake-transcript/{tid}/docx")
    assert code == 404


def test_intake_transcript_docx_route_404_for_missing_id(server):
    base, _ = server
    code, _, _ = _get(f"{base}/intake-transcript/999999/docx")
    assert code == 404
